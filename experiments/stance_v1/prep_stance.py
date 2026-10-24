"""
Siapkan input label sikap (stance) v1.

Output di experiments/stance_v1/:
  - stance_input.csv: 2.288 komentar relevan, diurutkan per video lalu per thread, dengan judul video
    dan teks komentar induk sebagai konteks.
  - validasi_sikap_zillan.xlsx: 100 komentar acak (seed 42) untuk validasi buta, kolom label kosong.
  - validasi_sikap_ids.csv: commentId sampel validasi.

Penggunaan:
  python experiments/stance_v1/prep_stance.py
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT_DIR = Path("experiments/stance_v1")
SEED = 42
N_VALID = 100


def main():
    rel = pd.read_csv("data/processed/comments_relevance_v1.csv", encoding="utf-8-sig", dtype={"commentId": str})
    rel = rel[rel["relevant"] == "y"]
    raw = pd.read_csv("data/20260930/comments.csv", encoding="utf-8-sig", dtype={"commentId": str, "parentId": str})
    videos = pd.read_csv("data/20260930/videos.csv", encoding="utf-8-sig")

    parent_text = raw.set_index("commentId")["textOriginal"]
    df = rel[["commentId", "videoId", "case_id", "org_type", "video_type", "isReply", "text_embed"]].merge(
        raw[["commentId", "parentId", "publishedAt"]], on="commentId", how="left")
    df["title"] = df["videoId"].map(videos.set_index("video_id")["title"])
    df["parent_text"] = df["parentId"].map(parent_text).fillna("")
    df["thread"] = df["parentId"].fillna(df["commentId"])
    df["thread_start"] = df.groupby("thread")["publishedAt"].transform("min")
    df = df.sort_values(["case_id", "videoId", "thread_start", "thread", "isReply", "publishedAt"])
    df.insert(0, "idx", range(1, len(df) + 1))
    cols = ["idx", "commentId", "case_id", "videoId", "title", "org_type", "video_type", "isReply",
            "parent_text", "text_embed"]
    df[cols].to_csv(OUT_DIR / "stance_input.csv", index=False, encoding="utf-8-sig")

    sample = df.sample(N_VALID, random_state=SEED).reset_index(drop=True)
    sample[["commentId"]].to_csv(OUT_DIR / "validasi_sikap_ids.csv", index=False)
    write_sheet(sample, OUT_DIR / "validasi_sikap_zillan.xlsx")
    print(f"{len(df)} komentar, {N_VALID} sampel validasi")


def write_sheet(sample: pd.DataFrame, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Label"
    header = ["no", "kasus", "judul video", "komentar induk (kalau balasan)", "komentar", "sikap", "catatan"]
    ws.append(header)
    for i, r in sample.iterrows():
        ws.append([i + 1, r["case_id"], r["title"], r["parent_text"], r["text_embed"], "", ""])
    widths = [5, 7, 30, 45, 60, 10, 30]
    for col, w in zip("ABCDEFG", widths):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in range(2, len(sample) + 2):
        ws[f"F{row}"].fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"tolak,dukung,lain"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F2:F{len(sample) + 1}")
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("Petunjuk")
    guide.column_dimensions["A"].width = 140
    for line in (OUT_DIR / "panduan_sikap.md").read_text(encoding="utf-8").splitlines():
        guide.append([line])
    for row in guide.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.move_sheet("Petunjuk", offset=-1)
    wb.active = 1
    wb.save(path)


if __name__ == "__main__":
    main()
