"""Build the adjudication sheet for the codebook v1 blind validation (rows where tema_utama differs)."""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

DIR = Path(__file__).parent
CODES = "KV,PK,HC,AK,KO,AU,EJ,LT"


def main():
    d = pd.read_csv(DIR / "validasi_tema_beda.csv").merge(pd.read_csv(DIR / "alasan_claude_validasi.csv"), on="no")
    path = DIR / "adjudikasi_tema.xlsx"
    if path.exists():  # never overwrite a sheet that may already be filled
        print(f"{path} sudah ada, tidak ditimpa.")
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Adjudikasi"
    ws.append(["no", "kasus", "komentar induk (kalau balasan)", "komentar", "validator", "catatan validator",
               "Claude", "alasan Claude", "penilaian ulang Claude", "keputusan tema_utama", "catatan"])
    fmt = lambda u, k: u if pd.isna(k) else f"{u} ({k})"
    for _, r in d.iterrows():
        ws.append([r["no"], r["case_id"], r["komentar induk (kalau balasan)"] if pd.notna(r["komentar induk (kalau balasan)"]) else "",
                   r["komentar"], fmt(r["tema_utama_human"], r["tema_kedua_human"]),
                   r["catatan"] if pd.notna(r["catatan"]) else "", fmt(r["tema_utama_claude"], r["tema_kedua_claude"]),
                   r["alasan_claude"], r["penilaian_ulang_claude"], "", ""])
    for col, w in zip("ABCDEFGHIJK", [5, 7, 40, 55, 11, 30, 11, 40, 40, 12, 30]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in range(2, len(d) + 2):
        ws[f"J{row}"].fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1=f'"{CODES}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J2:J{len(d) + 1}")
    ws.freeze_panes = "A2"
    wb.save(path)
    print(f"{len(d)} rows to adjudicate -> {path}")


if __name__ == "__main__":
    main()
