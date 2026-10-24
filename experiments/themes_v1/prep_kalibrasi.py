"""
Siapkan lembar kalibrasi codebook tema: 50 komentar pilot (pilot_codebook_ids.csv, seed 7)
untuk dikode manusia tanpa melihat kode Claude (pilot_claude.csv).

Output: experiments/themes_v1/kalibrasi_tema.xlsx (sheet Petunjuk = codebook, sheet Kode = lembar isian).
Nomor topik BERTopic sengaja tidak ditampilkan supaya pengode tidak terarah oleh topik.

Penggunaan:
  python experiments/themes_v1/prep_kalibrasi.py
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT_DIR = Path("experiments/themes_v1")
CODES = ["KV", "PK", "HC", "AK", "KO", "AU", "EJ", "LT"]  # codebook v1 (PB removed)


def main():
    ids = pd.read_csv(OUT_DIR / "pilot_codebook_ids.csv", dtype={"commentId": str})
    stance = pd.read_csv("experiments/stance_v1/stance_input.csv", dtype={"commentId": str})
    sample = ids.merge(stance, on=["commentId", "idx"], how="left")
    assert len(sample) == 50 and sample["text_embed"].notna().all()
    sample["parent_text"] = sample["parent_text"].fillna("")
    write_sheet(sample, OUT_DIR / "kalibrasi_tema.xlsx")
    print(f"{len(sample)} komentar ditulis ke {OUT_DIR / 'kalibrasi_tema.xlsx'}")


def write_sheet(sample: pd.DataFrame, path: Path, title: str = "Kalibrasi codebook tema (50 komentar)"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Kode"
    ws.append(["no", "kasus", "judul video", "komentar induk (kalau balasan)", "komentar",
               "tema_utama", "tema_kedua", "catatan"])
    for i, r in sample.iterrows():
        ws.append([i + 1, r["case_id"], r["title"], r["parent_text"], r["text_embed"], "", "", ""])
    for col, w in zip("ABCDEFGH", [5, 7, 30, 45, 60, 12, 12, 35]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    last = len(sample) + 1
    for row in range(2, last + 1):
        ws[f"F{row}"].fill = PatternFill("solid", fgColor="FFF2CC")
        ws[f"G{row}"].fill = PatternFill("solid", fgColor="FFF9E6")
    dv = DataValidation(type="list", formula1='"' + ",".join(CODES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F2:G{last}")
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("Petunjuk")
    guide.column_dimensions["A"].width = 140
    intro = [
        f"# {title}",
        "",
        "Isi sheet 'Kode': tema_utama wajib (satu kode), tema_kedua opsional. Pakai dropdown.",
        "Semua komentar sudah berlabel sikap 'tolak'; yang dikode adalah ALASAN keberatannya.",
        "Kerjakan sendiri tanpa melihat kode Claude. Kalau ragu, tetap pilih satu kode dan tulis keraguanmu di 'catatan'.",
        "Komentar yang menurutmu sebenarnya bukan keberatan: pilih kode terdekat dan tulis 'bukan tolak' di catatan.",
        "",
    ]
    for line in intro + (OUT_DIR / "codebook_tema.md").read_text(encoding="utf-8").splitlines():
        guide.append([line])
    for row in guide.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.move_sheet("Petunjuk", offset=-1)
    wb.active = 1
    wb.save(path)


if __name__ == "__main__":
    main()
